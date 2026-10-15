# -*- coding: utf-8 -*-
"""克隆学校《Java程序设计》期末考试报告单模板，填入九组论坛内容。"""
import shutil
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

SRC = r"C:/Users/ac010/Desktop/专升本文件/期末报告汇总/2025-2026学年第二学期---Java程序设计实践考核报告单张绍轩.docx"
OUT = r"C:/Users/ac010/WorkBuddy/2026-10-01-14-20-56/jiuzu-forum/九组论坛期末考试报告单.docx"

shutil.copy(SRC, OUT)
d = Document(OUT)
t = d.tables[0]


def set_cell(cell, text, size=12):
    """清空单元格并用指定字体重建多行段落（宋体/Time New Roman）。"""
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    for line in text.split("\n"):
        p = cell.add_paragraph()
        p.text = line
        if not p.runs:
            continue
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.name = "Times New Roman"
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


# ---- 个人信息（依据桌面文件名推断，教师待填）----
t.cell(0, 1).text = "25级数据科学与大数据技术（专升本）2班"
t.cell(0, 6).text = ""                       # 成绩留空
t.cell(1, 1).text = "2025091120239"
t.cell(1, 4).text = "张绍轩"
t.cell(1, 6).text = "【教师姓名】"

# ---- 项目演示：运行效果截图（改为九组论坛图1-7）----
set_cell(t.cell(4, 2),
"""1.项目相关运行效果截图
系统主界面
图1 系统运行主界面图（首页 / 板块列表）

基础操作界面
图2 用户登录 / 注册运行效果图
图3 发帖运行效果图
图4 帖子详情与回帖运行效果图
图5 点赞 / 搜索运行效果图""")

set_cell(t.cell(5, 2),
"""后台与 AI 界面
图6 管理员后台运行效果图
图7 AI 助手（/ai）与发帖“AI 帮写”运行效果图""")

# ---- 项目设计与实现过程及课程知识应用总结 ----
set_cell(t.cell(7, 2),
"""项目设计与实现过程及课程知识应用总结
项目设计与实现过程陈述
需求分析：
    九组论坛是面向课程考核与日常交流的校园在线论坛系统。用户需要浏览板块与帖子、发帖回帖、点赞收藏、搜索，以及管理员后台（用户 / 帖子 / 板块 / 回复管理）。系统需满足考核要求的“增删改查、图形化界面、数据库操作”。

方案设计：
    采用 MVC 架构。Model 层用 JavaBean（User / Board / Post / Reply）封装数据，DAO 层用 JDBC 直连 SQLite；View 层用 JSP + EL + JSTL 渲染并配合 CSS 实现响应式界面；Controller 层用 Servlet（BaseServlet 统一分发）处理请求。EncodingFilter 统一编码与缓存策略，AppListener 负责建库与种子数据初始化。

功能开发：
    实现用户注册登录（Session 鉴权）、帖子分页列表与详情、发帖回帖、点赞、搜索；管理员后台含数据统计、板块增删改（删除板块时可把帖子合并到其他板块）、回复管理。特色功能：集成 DeepSeek 大模型的 AI 助手（/ai）与发帖 / 回复“AI 帮写”，使用 JDK 原生 HttpURLConnection 调用，零第三方依赖。

测试优化：
    本地 JDK8 + Tomcat 8.5 编译部署，覆盖各功能路径；修复 JSTL 标签嵌套、浏览器缓存等问题；部署于腾讯云服务器并配置 systemd 自启、数据库备份与还原流程。

（2）课程知识应用总结
    本项目综合运用：① Servlet / JSP 实现请求分发与动态页面；② EL / JSTL 简化视图层展示；③ JavaBean 封装业务实体；④ JDBC 完成数据库增删改查；⑤ MVC 分层解耦；⑥ Filter / Listener 解决编码与初始化；⑦ Session 实现登录态；⑧ 原生 HTTP 调用第三方 AI 接口。完整覆盖课程核心知识点。""")

# ---- 项目答辩：问题一 / 二及答案 ----
set_cell(t.cell(9, 2),
"""问题一： 你的项目采用了什么架构？为什么这样设计？
答案： 采用 MVC（模型—视图—控制器）架构。Model 用 JavaBean + DAO 封装数据与数据库访问，View 用 JSP + EL + JSTL 渲染，Controller 用 Servlet 分发请求。分层使代码结构清晰、便于维护扩展，符合课程对“代码规范、逻辑清晰”的要求。

问题二： 你是如何接入 AI 功能的？是否引入额外框架？
答案： 通过 JDK 原生的 HttpURLConnection 直接调用 DeepSeek 开放接口（/v1/chat/completions），手工拼装与解析 JSON，未引入任何第三方 HTTP 或 AI 依赖。既贴合 Java Web 考试技术栈，又实现了 AI 助手与发帖 / 回复帮写。""")

d.save(OUT)
print("SAVED:", OUT)
